"""Tests for the evidence pipeline (task A1): extraction, indexing, structuring and
the processing status. Gemini is never called: a fake client stands in for it
through app.services.ai.client.get_client."""

import io
import json
from datetime import UTC, datetime
from types import SimpleNamespace

import docx
import pymupdf
import pytest
from google.genai import types
from sqlalchemy import select

from app.core.storage import save_bytes
from app.models import (
    AiRun,
    CaseObject,
    Conflict,
    Event,
    Evidence,
    EvidenceChunk,
    EvidenceLink,
    Location,
    Observation,
    Person,
    Relationship,
)
from app.services.ai import client as ai_client
from app.services.ai import gemini as gemini_module
from app.services.ai.extract import NO_TEXT, TRANSCRIBE_PROMPT
from app.services.ai.passages import ChunkFinder, normalize
from app.services.ai.pipeline import process_evidence
from app.services.ai.structure import parse_time
from tests.factories import make_case, make_event, make_evidence, make_object, make_observation, make_person


class FakeGemini:
    """Answers like google-genai's Client: structured JSON for calls with a
    response schema, OCR text or a transcript otherwise."""

    def __init__(self):
        self.calls: list[SimpleNamespace] = []
        self.structured: dict | str = {"summary": ""}
        self.ocr = "Text read from the image."
        self.transcript = "[00:00] Speaker 1: I saw the van move off."
        self.error: Exception | None = None
        self.uploads: list[SimpleNamespace] = []
        self.deleted: list[str] = []
        self.models = SimpleNamespace(generate_content=self._generate, embed_content=self._embed)
        self.files = SimpleNamespace(upload=self._upload, get=self._get, delete=self._delete)

    def _generate(self, *, model, contents, config=None):
        call = SimpleNamespace(model=model, contents=contents, config=config)
        self.calls.append(call)
        if config is not None and config.response_json_schema is not None:
            if self.error is not None:
                raise self.error
            structured = self.structured
            return SimpleNamespace(text=structured if isinstance(structured, str) else json.dumps(structured))
        return SimpleNamespace(text=self.transcript if TRANSCRIBE_PROMPT in contents else self.ocr)

    def _embed(self, *, model, contents, config=None):
        # For the search indexing that runs inside the pipeline once A2 lands.
        items = contents if isinstance(contents, list) else [contents]
        return SimpleNamespace(embeddings=[SimpleNamespace(values=[0.5] * 8) for _ in items])

    def _upload(self, *, file, config=None):
        uploaded = SimpleNamespace(name=f"files/{len(self.uploads) + 1}", state=types.FileState.PROCESSING)
        self.uploads.append(SimpleNamespace(file=file, config=config, uploaded=uploaded))
        return uploaded

    def _get(self, *, name, config=None):
        return SimpleNamespace(name=name, state=types.FileState.ACTIVE)

    def _delete(self, *, name, config=None):
        self.deleted.append(name)

    def structuring_calls(self):
        return [c for c in self.calls if c.config is not None and c.config.response_json_schema is not None]

    def media_calls(self):
        return [c for c in self.calls if c not in self.structuring_calls()]


@pytest.fixture
def gemini(monkeypatch):
    fake = FakeGemini()
    monkeypatch.setattr(ai_client, "get_client", lambda: fake)
    return fake


def stored_evidence(db, case, name: str, data: bytes, mime: str, **fields) -> Evidence:
    stored = save_bytes(case.id, name, data)
    values = {
        "title": name,
        "mime_type": mime,
        "file_path": stored.file_path,
        "size_bytes": stored.size_bytes,
        "sha256": stored.sha256,
        "processing_status": "queued",
    }
    return make_evidence(db, case, **(values | fields))


def pdf_bytes(*pages: str | None) -> bytes:
    """A PDF with one page per argument; None makes a page with no text layer."""
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
        else:
            page.draw_rect(pymupdf.Rect(72, 72, 300, 200), color=(0, 0, 0), fill=(0.5, 0.5, 0.5))
    return document.tobytes()


def docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("Statement of Anita Desai.")
    document.add_paragraph("")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Signal"
    table.rows[0].cells[1].text = "Amber"
    document.add_paragraph("The van moved off early.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def reload(db, evidence: Evidence) -> Evidence:
    db.expire_all()
    return db.get(Evidence, evidence.id)


def chunks_of(db, evidence: Evidence) -> list[EvidenceChunk]:
    query = select(EvidenceChunk).where(EvidenceChunk.evidence_id == evidence.id).order_by(EvidenceChunk.ordinal)
    return list(db.scalars(query))


def runs_of(db, task: str) -> list[AiRun]:
    return list(db.scalars(select(AiRun).where(AiRun.task == task).order_by(AiRun.id)))


def ai_rows(db, model) -> list:
    return list(db.scalars(select(model).where(model.origin == "ai").order_by(model.id)))


# ---------------------------------------------------------------------------
# Status and local extraction (no GEMINI_API_KEY)
# ---------------------------------------------------------------------------


def test_text_file_without_key_is_extracted_indexed_and_done(db):
    case = make_case(db)
    evidence = stored_evidence(
        db, case, "statement.txt", b"First paragraph.\r\n\r\nSecond paragraph.\r\n", "text/plain"
    )

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.processed_at is not None
    assert evidence.summary is None
    assert evidence.extracted_text == "First paragraph.\n\nSecond paragraph."
    chunks = chunks_of(db, evidence)
    assert normalize(" ".join(c.text for c in chunks)) == "first paragraph second paragraph"
    assert all(c.page is None for c in chunks)
    assert db.scalars(select(AiRun)).all() == []
    assert ai_rows(db, Observation) == []


def test_pdf_pages_keep_their_page_numbers(db):
    case = make_case(db)
    evidence = stored_evidence(
        db, case, "report.pdf", pdf_bytes("Page one text.", None, "Page three text."), "application/pdf"
    )

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.extracted_text == "Page one text.\n\nPage three text."
    assert [(c.page, c.text) for c in chunks_of(db, evidence)] == [(1, "Page one text."), (3, "Page three text.")]


def test_word_paragraphs_and_tables_are_read_in_order(db):
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    case = make_case(db)
    evidence = stored_evidence(db, case, "statement.docx", docx_bytes(), mime)

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.extracted_text == "Statement of Anita Desai.\n\nSignal | Amber\n\nThe van moved off early."


def test_type_is_guessed_from_the_file_name_when_the_upload_gives_none(db):
    case = make_case(db)
    evidence = stored_evidence(db, case, "notes.md", b"Scene notes.", "application/octet-stream")

    process_evidence(evidence.id)

    assert reload(db, evidence).extracted_text == "Scene notes."


def test_media_and_unknown_files_finish_without_text_when_ai_is_off(db):
    case = make_case(db)
    photo = stored_evidence(db, case, "photo.jpg", b"\xff\xd8 not really a jpeg", "image/jpeg")
    archive = stored_evidence(db, case, "dump.zip", b"PK", "application/zip")

    process_evidence(photo.id)
    process_evidence(archive.id)

    for evidence in (photo, archive):
        evidence = reload(db, evidence)
        assert evidence.processing_status == "done"
        assert evidence.extracted_text is None
        assert chunks_of(db, evidence) == []


def test_missing_file_marks_the_evidence_failed(db):
    case = make_case(db)
    evidence = make_evidence(db, case, file_path="1/missing.pdf", processing_status="queued")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "failed"
    assert evidence.summary == "Processing failed: The stored file is missing"
    assert evidence.processed_at is not None


def test_unknown_evidence_is_ignored():
    process_evidence(9999)


# ---------------------------------------------------------------------------
# OCR and transcription
# ---------------------------------------------------------------------------


def test_pdf_page_without_text_layer_is_read_by_ocr(db, gemini):
    gemini.ocr = "Handwritten note: van moved early."
    case = make_case(db)
    evidence = stored_evidence(db, case, "scan.pdf", pdf_bytes("Typed page.", None), "application/pdf")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.extracted_text == "Typed page.\n\nHandwritten note: van moved early."
    assert [(c.page, c.text) for c in chunks_of(db, evidence)] == [
        (1, "Typed page."),
        (2, "Handwritten note: van moved early."),
    ]
    [call] = gemini.media_calls()
    assert call.contents[0].inline_data.mime_type == "image/png"
    [run] = runs_of(db, "ocr")
    assert (run.evidence_id, run.case_id, run.model) == (evidence.id, case.id, "gemini-3.5-flash")
    assert run.output == {"page": 2, "text": "Handwritten note: van moved early."}


def test_image_text_is_read_by_ocr(db, gemini):
    case = make_case(db)
    evidence = stored_evidence(db, case, "sign.png", b"png bytes", "image/png")

    process_evidence(evidence.id)

    assert reload(db, evidence).extracted_text == "Text read from the image."
    [call] = gemini.media_calls()
    assert call.contents[0].inline_data.data == b"png bytes"
    assert call.contents[0].inline_data.mime_type == "image/png"


def test_image_without_text_has_no_extracted_text(db, gemini):
    gemini.ocr = NO_TEXT
    case = make_case(db)
    evidence = stored_evidence(db, case, "skid-marks.jpg", b"jpeg bytes", "image/jpeg")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.extracted_text is None
    assert len(runs_of(db, "ocr")) == 1
    assert gemini.structuring_calls() == []


def test_image_types_gemini_cannot_read_are_sent_as_png(db, gemini):
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8), False)
    pixmap.clear_with(255)
    case = make_case(db)
    evidence = stored_evidence(db, case, "scan.pnm", pixmap.tobytes("pnm"), "image/x-portable-anymap")

    process_evidence(evidence.id)

    [call] = gemini.media_calls()
    assert call.contents[0].inline_data.mime_type == "image/png"
    assert call.contents[0].inline_data.data.startswith(b"\x89PNG")


def test_audio_is_transcribed(db, gemini):
    case = make_case(db)
    evidence = stored_evidence(db, case, "call.mp3", b"mp3 bytes", "audio/mpeg")

    process_evidence(evidence.id)

    assert reload(db, evidence).extracted_text == "[00:00] Speaker 1: I saw the van move off."
    [call] = gemini.media_calls()
    assert call.contents[0].inline_data.mime_type == "audio/mpeg"
    assert call.contents[1] == TRANSCRIBE_PROMPT
    [run] = runs_of(db, "transcribe")
    assert run.output == {"text": "[00:00] Speaker 1: I saw the van move off."}


def test_large_video_goes_through_the_files_api(db, gemini, monkeypatch):
    monkeypatch.setattr(gemini_module, "INLINE_LIMIT", 4)
    monkeypatch.setattr(gemini_module.time, "sleep", lambda seconds: None)
    case = make_case(db)
    evidence = stored_evidence(db, case, "dashcam.mp4", b"a longer video", "video/mp4")

    process_evidence(evidence.id)

    assert reload(db, evidence).extracted_text == "[00:00] Speaker 1: I saw the van move off."
    [upload] = gemini.uploads
    assert upload.config.mime_type == "video/mp4"
    [call] = gemini.media_calls()
    assert call.contents[0].name == "files/1"
    assert call.contents[0].state == types.FileState.ACTIVE
    assert gemini.deleted == ["files/1"]


# ---------------------------------------------------------------------------
# Structuring
# ---------------------------------------------------------------------------

STATEMENT = """Statement of Anita Desai, recorded on 4 October 2026.

I was standing at the bus stop at Ring Road and 5th Avenue at about 9:40 pm. The white hatchback, KA 01 AB 1234, entered the junction on amber.

The blue delivery van moved off before its light turned green. Mr. Rohan Kapoor was driving the hatchback."""

STRUCTURED = {
    "summary": "A witness at the bus stop says the hatchback entered on amber and the van moved off early.",
    "persons": [
        {"ref": "p1", "name": "Anita Desai", "role": "witness", "passage": "Statement of Anita Desai"},
        {
            "ref": "p2",
            "name": "Mr. Rohan Kapoor",
            "role": "victim",
            "passage": "Mr. Rohan Kapoor was driving the hatchback.",
        },
    ],
    "objects": [
        {
            "ref": "o1",
            "kind": "vehicle",
            "label": "White hatchback",
            "plate": "KA 01 AB 1234",
            "passage": "The white hatchback, KA 01 AB 1234, entered the junction on amber.",
        },
        {
            "ref": "o2",
            "kind": "vehicle",
            "label": "Blue delivery van",
            "colour": "blue",
            "passage": "The blue delivery van moved off before its light turned green.",
        },
    ],
    "locations": [
        {
            "ref": "l1",
            "kind": "incident_site",
            "address": "Ring Road and 5th Avenue",
            "passage": "I was standing at the bus stop at Ring Road and 5th Avenue",
        }
    ],
    "events": [
        {
            "ref": "e1",
            "category": "movement",
            "description": "The van moves off before its light turns green.",
            "time_start": "2026-10-03T21:40:00+05:30",
            "time_end": None,
            "time_precision": "approximate",
            "certainty": "hypothesis",
            "location_ref": "l1",
            "note": "Time is the witness's estimate.",
            "passage": "The blue delivery van moved off before its light turned green.",
        }
    ],
    "relationships": [
        {
            "from_ref": "p2",
            "to_ref": "o1",
            "kind": "drove",
            "confidence": 1.4,
            "passage": "Mr. Rohan Kapoor was driving the hatchback.",
        },
        {
            "from_ref": "evidence",
            "to_ref": "p1",
            "kind": "mentions",
            "label": "Witness",
            "passage": "Statement of Anita Desai",
        },
        {"from_ref": "p9", "to_ref": "o1", "kind": "owns", "passage": "Statement of Anita Desai"},
        {"from_ref": "e1", "to_ref": "o2", "kind": "involved_in", "passage": "The blue delivery van moved off"},
    ],
    "observations": [
        {
            "subject_ref": "o1",
            "statement": "The witness says the hatchback entered the junction on amber.",
            "observed_at": "2026-10-03T21:40:00+05:30",
            "passage": "The white hatchback, KA 01 AB 1234, entered the junction on amber.",
        },
        {
            "subject_ref": "o2",
            "statement": "The witness says the van moved off before its light turned green.",
            "observed_at": None,
            "passage": "The blue delivery van moved off before its light turned green.",
        },
        {
            "subject_ref": None,
            "statement": "The statement was recorded the next day.",
            "passage": "recorded on 4 October 2026",
        },
    ],
}


def statement_evidence(db, case, name="witness-statement.txt") -> Evidence:
    return stored_evidence(db, case, name, STATEMENT.encode(), "text/plain", type="statement")


def test_structuring_saves_pending_proposals_linked_to_their_passages(db, gemini):
    gemini.structured = STRUCTURED
    case = make_case(db)
    rohan = make_person(db, case, name="Rohan Kapoor", role="victim")
    hatchback = make_object(db, case, label="White hatchback", details={"plate": "KA-01-AB-1234"})
    evidence = statement_evidence(db, case)

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.summary == STRUCTURED["summary"]

    # The raw output is logged.
    [run] = runs_of(db, "structure")
    assert (run.evidence_id, run.case_id, run.output) == (evidence.id, case.id, STRUCTURED)

    # The prompt names the known records, so the model can reuse them.
    [call] = gemini.structuring_calls()
    prompt = call.contents[0]
    assert "Rohan Kapoor (victim)" in prompt and "plate KA-01-AB-1234" in prompt and STATEMENT in prompt

    # Rohan and the hatchback are matched, not duplicated; the others are proposals.
    assert [p.name for p in db.scalars(select(Person).order_by(Person.id))] == ["Rohan Kapoor", "Anita Desai"]
    db.refresh(rohan)
    assert (rohan.origin, rohan.review_status) == ("investigator", "confirmed")
    [anita] = ai_rows(db, Person)
    assert (anita.role, anita.review_status) == ("witness", "pending")
    [van] = ai_rows(db, CaseObject)
    assert (van.label, van.details, van.review_status) == ("Blue delivery van", {"colour": "blue"}, "pending")
    [junction] = ai_rows(db, Location)
    assert (junction.kind, junction.address) == ("incident_site", "Ring Road and 5th Avenue")

    [event] = ai_rows(db, Event)
    assert event.location_id == junction.id
    assert event.time_start == datetime(2026, 10, 3, 16, 10, tzinfo=UTC)
    assert (event.time_precision, event.certainty, event.ai_note) == (
        "approximate",
        "hypothesis",
        "Time is the witness's estimate.",
    )

    # Relationships with unknown refs or non-graph ends are skipped.
    relationships = ai_rows(db, Relationship)
    assert [(r.from_type, r.from_id, r.to_type, r.to_id, r.kind, r.label, r.confidence) for r in relationships] == [
        ("person", rohan.id, "object", hatchback.id, "drove", "Drove", 1.0),
        ("evidence", evidence.id, "person", anita.id, "mentions", "Witness", None),
    ]

    observations = ai_rows(db, Observation)
    assert [(o.subject_type, o.subject_id) for o in observations] == [
        ("object", hatchback.id),
        ("object", van.id),
        ("case", case.id),
    ]
    assert observations[0].observed_at == datetime(2026, 10, 3, 16, 10, tzinfo=UTC)
    assert observations[1].observed_at is None
    assert all(o.review_status == "pending" and o.author_id is None for o in observations)

    # Every proposal and matched record links to its passage and the chunk holding it.
    links = list(db.scalars(select(EvidenceLink).order_by(EvidenceLink.id)))
    targets = {(link.target_type, link.target_id) for link in links}
    assert targets == {
        ("location", junction.id),
        ("person", anita.id),
        ("person", rohan.id),
        ("object", hatchback.id),
        ("object", van.id),
        ("event", event.id),
        *(("relationship", r.id) for r in relationships),
        *(("observation", o.id) for o in observations),
    }
    chunks = {c.id: c for c in chunks_of(db, evidence)}
    for link in links:
        assert link.evidence_id == evidence.id
        assert normalize(link.passage) in normalize(chunks[link.chunk_id].text)


def test_reprocessing_replaces_pending_proposals_and_keeps_reviewed_ones(db, gemini):
    gemini.structured = STRUCTURED
    case = make_case(db)
    evidence = statement_evidence(db, case)
    process_evidence(evidence.id)

    # An investigator confirms Anita and links an event of their own to a
    # passage, and an open conflict cites one of the pending observations.
    anita = db.scalar(select(Person).where(Person.name == "Anita Desai"))
    anita.review_status = "confirmed"
    own_event = make_event(db, case, description="Witness interviewed.")
    first_chunk = chunks_of(db, evidence)[0]
    db.add(
        EvidenceLink(
            evidence_id=evidence.id,
            chunk_id=first_chunk.id,
            target_type="event",
            target_id=own_event.id,
            passage="Statement of Anita Desai",
        )
    )
    first_observations = [o.id for o in ai_rows(db, Observation)]
    other = make_observation(db, case, statement="The light was red.")
    db.add(
        Conflict(
            case_id=case.id,
            observation_ids=[other.id, first_observations[0]],
            description="Signal colour",
            status="open",
        )
    )
    db.commit()

    process_evidence(evidence.id)

    assert reload(db, evidence).processing_status == "done"
    persons = list(db.scalars(select(Person).order_by(Person.id)))
    assert [(p.name, p.review_status) for p in persons] == [
        ("Anita Desai", "confirmed"),
        ("Mr. Rohan Kapoor", "pending"),
    ]
    assert len(ai_rows(db, CaseObject)) == 2
    assert len(ai_rows(db, Location)) == 1
    assert len(ai_rows(db, Event)) == 1
    assert len(ai_rows(db, Relationship)) == 2
    observations = ai_rows(db, Observation)
    assert len(observations) == 3
    assert not set(first_observations) & {o.id for o in observations}
    assert db.scalars(select(Conflict)).all() == []
    assert db.get(Observation, other.id) is not None

    # Re-indexing emptied every link's chunk_id; each now points at a new chunk,
    # including the investigator's own link.
    chunk_ids = {c.id for c in chunks_of(db, evidence)}
    links = list(db.scalars(select(EvidenceLink)))
    assert all(link.chunk_id in chunk_ids for link in links)
    assert any(link.target_type == "event" and link.target_id == own_event.id for link in links)
    assert sum(1 for link in links if (link.target_type, link.target_id) == ("person", anita.id)) == 1


def test_reprocessing_keeps_pending_proposals_other_evidence_supports(db, gemini):
    gemini.structured = {
        "summary": "Names a witness.",
        "persons": [{"ref": "p1", "name": "Anita Desai", "role": "witness", "passage": "Statement of Anita Desai"}],
    }
    case = make_case(db)
    first = statement_evidence(db, case, "first.txt")
    second = statement_evidence(db, case, "second.txt")
    process_evidence(first.id)
    process_evidence(second.id)
    [anita] = ai_rows(db, Person)

    process_evidence(first.id)

    assert [p.id for p in ai_rows(db, Person)] == [anita.id]
    linked_from = sorted(db.scalars(select(EvidenceLink.evidence_id).where(EvidenceLink.target_type == "person")))
    assert linked_from == [first.id, second.id]


def test_gemini_error_marks_failed_and_keeps_the_text(db, gemini):
    gemini.error = RuntimeError("429 RESOURCE_EXHAUSTED")
    case = make_case(db)
    evidence = statement_evidence(db, case)

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "failed"
    assert evidence.summary == "Processing failed: 429 RESOURCE_EXHAUSTED"
    assert evidence.processed_at is not None
    assert evidence.extracted_text == STATEMENT
    assert chunks_of(db, evidence) != []


def test_invalid_json_is_logged_and_marks_failed(db, gemini):
    gemini.structured = "Sorry, I can't help with that."
    case = make_case(db)
    evidence = statement_evidence(db, case)

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "failed"
    assert evidence.summary == "Processing failed: Gemini did not return valid JSON"
    [run] = runs_of(db, "structure")
    assert run.output == "Sorry, I can't help with that."
    assert ai_rows(db, Person) == []


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def test_parse_time():
    assert parse_time("2026-10-03T21:40:00+05:30") == datetime(2026, 10, 3, 16, 10, tzinfo=UTC)
    assert parse_time("2026-10-03T16:10:00Z") == datetime(2026, 10, 3, 16, 10, tzinfo=UTC)
    assert parse_time("2026-10-03T16:10") == datetime(2026, 10, 3, 16, 10, tzinfo=UTC)
    assert parse_time("about 9:40 pm") is None
    assert parse_time(None) is None


def test_chunk_finder_matches_verbatim_then_by_overlap():
    first = EvidenceChunk(id=1, text="The van moved off\nbefore its light turned green.")
    second = EvidenceChunk(id=2, text="The hatchback entered on amber.")
    finder = ChunkFinder([first, second])
    assert finder.find("the van moved off before its light turned GREEN") is first
    assert finder.find("The hatchback entered the junction on amber") is second
    assert finder.find("Nobody saw the pedestrian.") is None
    assert finder.find("") is None
